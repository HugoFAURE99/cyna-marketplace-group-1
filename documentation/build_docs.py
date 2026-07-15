"""Génère les six documents Word CYNA à partir des sources Markdown.

Le parseur couvre volontairement un sous-ensemble Markdown contrôlé afin de
garantir un rendu Word stable : titres, paragraphes, listes, tableaux, images,
citations, blocs de code et directives de section paysage.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE_DIR = ROOT / "sources"
TOC_DIR = ROOT / ".toc"

GROUP_NAME = "Groupe 1"
AUTHORS = ("Hugo FAURE", "Samuel KAKEZ", "Mathieu FRAUX", "Florian BELLOC")

OUTPUTS = {
    "01-sfd.md": "1 - SFD_Specification_fonctionnelle.docx",
    "02-dat.md": "2 - DAT_Architecture_technique.docx",
    "03-installation.md": "3 - Guide_installation.docx",
    "04-api.md": "4 - Documentation_API.docx",
    "05-securite.md": "5 - Plan_de_securite.docx",
    "06-maintenance.md": "6 - Plan_de_maintenance.docx",
}

INK = "1A1A24"
PURPLE = "302082"
PURPLE_2 = "4A339E"
PALE = "EEEAFB"
PALE_BLUE = "E8EEF5"
MUTED = "626273"
WHITE = "FFFFFF"


@dataclass
class Source:
    meta: dict[str, str]
    lines: list[str]


def parse_source(path: Path) -> Source:
    raw = path.read_text(encoding="utf-8").replace("\r\n", "\n")
    lines = raw.splitlines()
    meta: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1)
        for line in lines[1:end]:
            if ":" in line:
                key, value = line.split(":", 1)
                meta[key.strip()] = value.strip().strip('"')
        lines = lines[end + 1 :]
    return Source(meta=meta, lines=lines)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table, color: str = "C9C5DB") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def bookmark_name(text: str, index: int) -> str:
    slug = re.sub(r"[^A-Za-z0-9_]", "_", text)
    return f"h_{index}_{slug[:28]}"


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(
    paragraph,
    text: str,
    anchor: str,
    *,
    color=PURPLE,
    font_size_pt: float | None = None,
) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    if font_size_pt is not None:
        half_points = str(int(round(font_size_pt * 2)))
        size_node = OxmlElement("w:sz")
        size_node.set(qn("w:val"), half_points)
        r_pr.append(size_node)
        size_cs_node = OxmlElement("w:szCs")
        size_cs_node.set(qn("w:val"), half_points)
        r_pr.append(size_cs_node)
    run.append(r_pr)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


def configure_section(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(1.8)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Title", 26, PURPLE, 0, 10),
        ("Heading 1", 17, PURPLE, 18, 9),
        ("Heading 2", 13.5, PURPLE, 14, 7),
        ("Heading 3", 11.5, PURPLE_2, 10, 5),
        ("Heading 4", 10.5, PURPLE_2, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    if "CYNA Subtitle" not in styles:
        subtitle = styles.add_style("CYNA Subtitle", WD_STYLE_TYPE.PARAGRAPH)
        subtitle.base_style = normal
    subtitle = styles["CYNA Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(PURPLE_2)
    subtitle.paragraph_format.space_after = Pt(14)

    if "CYNA Caption" not in styles:
        caption = styles.add_style("CYNA Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.base_style = normal
    caption = styles["CYNA Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)

    if "CYNA Code" not in styles:
        code = styles.add_style("CYNA Code", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
    code = styles["CYNA Code"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.left_indent = Cm(0.45)
    code.paragraph_format.right_indent = Cm(0.25)
    code.paragraph_format.space_after = Pt(4)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.12


def add_inline_markdown(paragraph, text: str) -> None:
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`|\[.+?\]\(.+?\))")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(PURPLE_2)
        else:
            label, url = re.match(r"\[(.+?)\]\((.+?)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor.from_string(PURPLE_2)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def table_rows(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(doc: Document, rows: list[list[str]], landscape: bool) -> None:
    if not rows:
        return
    width_count = max(len(row) for row in rows)
    rows = [row + [""] * (width_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=width_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    available = 14600 if landscape else 9180
    weights = []
    for column in range(width_count):
        # A minimum weight keeps short header labels such as "Paramètre" from
        # being squeezed into unreadable one-letter lines in dense tables.
        longest = max(12, max(len(row[column]) for row in rows))
        weights.append(min(longest, 42))
    total = sum(weights)
    widths = [max(900, int(available * weight / total)) for weight in weights]
    scale = available / sum(widths)
    widths = [int(value * scale) for value in widths]

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, (row, values) in enumerate(zip(table.rows, rows)):
        for column, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell_width(cell, widths[column])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline_markdown(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(8.2 if width_count >= 5 else 8.8)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(PURPLE)
            if row_index == 0:
                shade(cell, PALE_BLUE)
            elif row_index % 2 == 0:
                shade(cell, "F8F7FC")
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def heading_inventory(lines: Iterable[str]) -> list[tuple[int, str, str]]:
    result = []
    index = 1
    for line in lines:
        match = re.match(r"^(#{1,4})\s+(.+?)\s*$", line)
        if match:
            level = len(match.group(1))
            title = match.group(2).strip()
            result.append((level, title, bookmark_name(title, index)))
            index += 1
    return result


def add_cover(doc: Document, meta: dict[str, str]) -> None:
    logo = ROOT / "diagrammes" / "cyna-logo.png"
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    if logo.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_shape = p_logo.add_run().add_picture(str(logo), width=Cm(4.6))
        logo_shape._inline.docPr.set("descr", "Logo CYNA")
        logo_shape._inline.docPr.set("title", "Logo CYNA")
    for _ in range(2):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(meta.get("title", "Documentation CYNA"))
    subtitle = doc.add_paragraph(style="CYNA Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(meta.get("subtitle", "Référence projet"))
    accent = doc.add_table(rows=1, cols=1)
    accent.alignment = WD_TABLE_ALIGNMENT.CENTER
    accent.autofit = False
    set_cell_width(accent.cell(0, 0), 5200)
    shade(accent.cell(0, 0), PURPLE)
    accent.cell(0, 0).text = ""
    set_repeat_table_header(accent.rows[0])
    doc.add_paragraph()
    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    labels = (
        ("Référence", meta.get("reference", "CYNA-DOC-002")),
        ("Version", meta.get("version", "2.0")),
        ("Date", meta.get("date", "15 juillet 2026")),
        ("Statut", meta.get("status", "Référence technique")),
    )
    for index, (label, value) in enumerate(labels):
        left, right = info.rows[index].cells
        set_cell_width(left, 1800)
        set_cell_width(right, 4200)
        set_cell_margins(left, 90, 130, 90, 130)
        set_cell_margins(right, 90, 130, 90, 130)
        left.text = label
        right.text = value
        left.paragraphs[0].runs[0].bold = True
        left.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(PURPLE)
        if index % 2 == 0:
            shade(left, PALE_BLUE)
            shade(right, "F8F7FC")
    set_repeat_table_header(info.rows[0])
    set_table_borders(info)
    audience = doc.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience.paragraph_format.space_before = Pt(18)
    run = audience.add_run(f"Audience : {meta.get('audience', 'équipe projet, exploitation et jury')}")
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(MUTED)
    team = doc.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team.paragraph_format.space_before = Pt(8)
    team.paragraph_format.space_after = Pt(0)
    team_run = team.add_run(f"{GROUP_NAME} - Équipe projet")
    team_run.bold = True
    team_run.font.color.rgb = RGBColor.from_string(PURPLE)
    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_after = Pt(0)
    authors_run = authors.add_run("  ·  ".join(AUTHORS))
    authors_run.font.size = Pt(9)
    authors_run.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()


def add_headers_and_footers(doc: Document, meta: dict[str, str]) -> None:
    for section in doc.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.clear()
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.text = f"CYNA  ·  {meta.get('reference', 'Documentation V2')}"
        for run in header.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"CYNA | {GROUP_NAME}  ·  ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        page_run = footer.add_run()
        page_run.font.size = Pt(8)
        add_field(page_run, "PAGE")


def add_toc(doc: Document, headings, page_map: dict[str, int]) -> None:
    p = doc.add_heading("Sommaire", level=1)
    add_bookmark(p, "toc", 9990)
    entries = [(level, title, anchor) for level, title, anchor in headings if level <= 3]
    compact = len(entries) > 36
    toc_font_size = 9.5 if compact else 10.5
    for level, title, anchor in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm((level - 1) * 0.45)
        p.paragraph_format.space_after = Pt(0 if compact else 2)
        p.paragraph_format.line_spacing = 1.0
        tab = p.paragraph_format.tab_stops
        tab.add_tab_stop(Cm(15.3), alignment=2, leader=1)
        add_internal_link(p, title, anchor, font_size_pt=toc_font_size)
        tab_run = p.add_run("\t")
        tab_run.font.size = Pt(toc_font_size)
        page = page_map.get(title)
        page_run = p.add_run(str(page) if page else "n/a")
        page_run.font.size = Pt(toc_font_size)
    doc.add_page_break()


def add_image(doc: Document, source_path: Path, caption: str, landscape: bool) -> None:
    image_path = source_path.resolve()
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    max_width = Cm(24.5 if landscape else 15.7)
    max_height = Cm(15.0 if landscape else 19.5)
    ratio = width_px / max(height_px, 1)
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(image_path), width=width, height=height)
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    cap = doc.add_paragraph(style="CYNA Caption")
    cap.add_run(caption)


def build_document(source_path: Path, output_path: Path) -> None:
    source = parse_source(source_path)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], landscape=False)
    props = doc.core_properties
    props.title = source.meta.get("title", output_path.stem)
    props.subject = source.meta.get("subtitle", "Documentation technique CYNA")
    props.author = f"{GROUP_NAME}: {', '.join(AUTHORS)}"
    props.last_modified_by = GROUP_NAME
    props.comments = "Généré depuis les sources Markdown versionnées."
    props.keywords = "CYNA, documentation, cybersécurité, SaaS"

    add_cover(doc, source.meta)
    headings = heading_inventory(source.lines)
    page_file = TOC_DIR / f"{source_path.stem}.json"
    page_map = json.loads(page_file.read_text(encoding="utf-8")) if page_file.exists() else {}
    add_toc(doc, headings, page_map)

    index = 0
    heading_index = 1
    in_code = False
    code_lines: list[str] = []
    landscape = False
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline_markdown(p, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer = []

    while index < len(source.lines):
        raw = source.lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                p = doc.add_paragraph(style="CYNA Code")
                for offset, code_line in enumerate(code_lines):
                    if offset:
                        p.add_run().add_break()
                    p.add_run(code_line)
                shade_code = OxmlElement("w:shd")
                shade_code.set(qn("w:fill"), "F3F1F8")
                p._p.get_or_add_pPr().append(shade_code)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped == "<!-- landscape-start -->":
            flush_paragraph()
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            configure_section(section, landscape=True)
            landscape = True
            index += 1
            continue
        if stripped == "<!-- landscape-end -->":
            flush_paragraph()
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            configure_section(section, landscape=False)
            landscape = False
            index += 1
            continue
        if stripped == "<!-- page-break -->":
            flush_paragraph()
            doc.add_page_break()
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.+?)\s*$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            p = doc.add_heading(title, level=level)
            add_bookmark(p, bookmark_name(title, heading_index), heading_index)
            heading_index += 1
            index += 1
            continue

        image = re.match(r"^!\[(.+?)\]\((.+?)\)\s*$", stripped)
        if image:
            flush_paragraph()
            add_image(doc, source_path.parent / image.group(2), image.group(1), landscape)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = table_rows(source.lines, index)
            add_table(doc, rows, landscape)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        number = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if bullet or number:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline_markdown(p, (bullet or number).group(1))
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.right_indent = Cm(0.25)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            shade_callout = OxmlElement("w:shd")
            shade_callout.set(qn("w:fill"), PALE)
            p._p.get_or_add_pPr().append(shade_callout)
            run = p.add_run(stripped[1:].strip())
            run.font.color.rgb = RGBColor.from_string(PURPLE)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
        else:
            paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    add_headers_and_footers(doc, source.meta)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", help="Nom d'une source Markdown, sinon toutes.")
    args = parser.parse_args()
    selected = [args.source] if args.source else list(OUTPUTS)
    for source_name in selected:
        if source_name not in OUTPUTS:
            raise SystemExit(f"Source inconnue : {source_name}")
        source_path = SOURCE_DIR / source_name
        output_path = ROOT / OUTPUTS[source_name]
        build_document(source_path, output_path)
        print(f"Généré : {output_path.name}")


if __name__ == "__main__":
    main()
